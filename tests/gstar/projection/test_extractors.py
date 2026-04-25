"""확장자별 텍스트 추출기 테스트."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

from gstar.projection.extractors import (
    combined_text,
    extract_directory,
    extract_file,
    supported_exts,
)


def test_supported_exts_covers_core_formats():
    exts = supported_exts()
    for e in [".md", ".txt", ".pdf", ".hwp", ".hwpx", ".docx", ".pptx", ".xlsx", ".json"]:
        assert e in exts, f"missing {e}"


def test_supported_exts_includes_code():
    exts = supported_exts()
    for e in [".py", ".ts", ".tsx", ".js", ".go", ".rs", ".sh"]:
        assert e in exts, f"missing {e}"


def test_extract_md(tmp_path: Path):
    f = tmp_path / "note.md"
    f.write_text("# 제목\n\n본문 한줄.", encoding="utf-8")
    r = extract_file(f)
    assert r.ok
    assert "제목" in r.text
    assert r.ext == ".md"


def test_extract_txt(tmp_path: Path):
    f = tmp_path / "t.txt"
    f.write_text("plain text content", encoding="utf-8")
    assert extract_file(f).ok


def test_extract_json_pretty_prints(tmp_path: Path):
    f = tmp_path / "form.json"
    f.write_text(json.dumps({"form_id": "abc", "title": "T"}), encoding="utf-8")
    r = extract_file(f)
    assert r.ok
    assert "form_id" in r.text
    assert "abc" in r.text


def test_extract_csv(tmp_path: Path):
    f = tmp_path / "data.csv"
    f.write_text("col1,col2\nval1,val2\nval3,val4\n", encoding="utf-8")
    r = extract_file(f)
    assert r.ok
    assert "val1" in r.text


def test_extract_cp949_fallback(tmp_path: Path):
    f = tmp_path / "legacy.md"
    f.write_bytes("한글 내용".encode("cp949"))
    r = extract_file(f)
    assert r.ok
    assert "한글" in r.text


def test_extract_nonexistent_file():
    r = extract_file(Path("/nonexistent/xyz.md"))
    assert not r.ok
    assert "error" in r.meta


def test_extract_unsupported_ext(tmp_path: Path):
    f = tmp_path / "a.bin"
    f.write_bytes(b"\x00\x01\x02")
    r = extract_file(f)
    assert not r.ok


def test_extract_hwpx_from_zip_xml(tmp_path: Path):
    hwpx = tmp_path / "sample.hwpx"
    xml_content = (
        """<?xml version="1.0" encoding="UTF-8"?>
<hp:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">
  <hp:p><hp:run><hp:t>첫 단락 내용</hp:t></hp:run></hp:p>
  <hp:p><hp:run><hp:t>둘째 단락 KPI AP@0.5 = 85%</hp:t></hp:run></hp:p>
</hp:sec>
"""
    ).encode("utf-8")
    with zipfile.ZipFile(hwpx, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("Contents/section0.xml", xml_content)
        zf.writestr("mimetype", "application/hwp+zip")
    r = extract_file(hwpx)
    assert r.ok, r.meta
    assert "첫 단락 내용" in r.text
    assert "KPI" in r.text


def test_extract_hwpx_namespace_agnostic(tmp_path: Path):
    hwpx = tmp_path / "plain.hwpx"
    xml_content = b"""<?xml version="1.0"?><sec><p><t>no ns</t></p></sec>"""
    with zipfile.ZipFile(hwpx, "w") as zf:
        zf.writestr("Contents/section0.xml", xml_content)
    r = extract_file(hwpx)
    assert "no ns" in r.text


def test_extract_hwpx_table_to_markdown(tmp_path: Path):
    """hp:tbl 이 마크다운 표로 변환되어 라벨-값 구조가 살아있어야 한다."""
    hwpx = tmp_path / "tbl.hwpx"
    xml = (
        """<?xml version="1.0" encoding="UTF-8"?>
<hp:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">
  <hp:p><hp:run><hp:t>본문 단락</hp:t></hp:run></hp:p>
  <hp:tbl>
    <hp:tr>
      <hp:tc><hp:p><hp:run><hp:t>과제명</hp:t></hp:run></hp:p></hp:tc>
      <hp:tc><hp:p><hp:run><hp:t>AI 반도체</hp:t></hp:run></hp:p></hp:tc>
    </hp:tr>
    <hp:tr>
      <hp:tc><hp:p><hp:run><hp:t>주관기관</hp:t></hp:run></hp:p></hp:tc>
      <hp:tc><hp:p><hp:run><hp:t>다겸</hp:t></hp:run></hp:p></hp:tc>
    </hp:tr>
  </hp:tbl>
</hp:sec>
"""
    ).encode("utf-8")
    with zipfile.ZipFile(hwpx, "w") as zf:
        zf.writestr("Contents/section0.xml", xml)
    r = extract_file(hwpx)
    assert r.ok, r.meta
    assert "본문 단락" in r.text
    assert "| 과제명 | AI 반도체 |" in r.text
    assert "| 주관기관 | 다겸 |" in r.text
    assert "| --- | --- |" in r.text


def test_extract_pptx(tmp_path: Path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    box = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "프로젝트 개요"

    s2 = prs.slides.add_slide(blank)
    tbl = s2.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(5), Inches(2)
    ).table
    tbl.cell(0, 0).text = "헤더1"
    tbl.cell(0, 1).text = "헤더2"
    tbl.cell(1, 0).text = "값1"
    tbl.cell(1, 1).text = "값2"

    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    r = extract_file(path)
    assert r.ok, r.meta
    assert "Slide 1" in r.text
    assert "프로젝트 개요" in r.text
    assert "Slide 2" in r.text
    assert "| 헤더1 | 헤더2 |" in r.text
    assert "| 값1 | 값2 |" in r.text


def test_extract_code_python(tmp_path: Path):
    f = tmp_path / "main.py"
    f.write_text("def hello():\n    return 'world'\n", encoding="utf-8")
    r = extract_file(f)
    assert r.ok
    assert r.ext == ".py"
    assert "hello" in r.text


def test_extract_code_typescript(tmp_path: Path):
    f = tmp_path / "app.ts"
    f.write_text("export const x: number = 42;\n", encoding="utf-8")
    r = extract_file(f)
    assert r.ok
    assert "x: number" in r.text


def test_extract_hwp_no_tooling_returns_empty(tmp_path: Path, monkeypatch):
    """hwp5html CLI 가 PATH·sys.executable 양쪽에 모두 없으면 빈 결과 반환."""
    import shutil
    import sys

    fake = tmp_path / "fake.hwp"
    fake.write_bytes(b"\xd0\xcf\x11\xe0")
    monkeypatch.setattr(shutil, "which", lambda _: None)
    monkeypatch.setattr(sys, "executable", str(tmp_path / "no_python"))
    r = extract_file(fake)
    assert r.ext == ".hwp"
    assert r.text == ""


def test_extract_hwpml_xml(tmp_path: Path):
    """`.hwp` 확장자에 HWPML XML 내용 — 정부 법제처가 export 하는 변종."""
    hwpml = tmp_path / "rule.hwp"
    xml = """<?xml version="1.0" encoding="utf-8"?>
<HWPML Version="2.1">
  <HEAD><DOCSUMMARY><TITLE>샘플 훈령</TITLE></DOCSUMMARY></HEAD>
  <BODY>
    <SECTION Id="0">
      <P><TEXT><CHAR>제1조(목적) 본 지침은 시범 목적이다.</CHAR></TEXT></P>
      <P><TEXT><CHAR>제2조(정의) 용어는 다음과 같다.</CHAR></TEXT></P>
      <TABLE>
        <ROW>
          <CELL><P><TEXT><CHAR>구분</CHAR></TEXT></P></CELL>
          <CELL><P><TEXT><CHAR>금액</CHAR></TEXT></P></CELL>
        </ROW>
        <ROW>
          <CELL><P><TEXT><CHAR>인건비</CHAR></TEXT></P></CELL>
          <CELL><P><TEXT><CHAR>1억원</CHAR></TEXT></P></CELL>
        </ROW>
      </TABLE>
    </SECTION>
  </BODY>
</HWPML>
"""
    hwpml.write_text(xml, encoding="utf-8")
    r = extract_file(hwpml)
    assert r.ok, r.meta
    assert "제1조(목적) 본 지침은 시범 목적이다." in r.text
    assert "제2조(정의)" in r.text
    assert "| 구분 | 금액 |" in r.text
    assert "| 인건비 | 1억원 |" in r.text
    assert "| --- | --- |" in r.text


def test_extract_hwp_unsupported_xml_root_errors(tmp_path: Path):
    """HWPML 도 OLE2 도 아닌 일반 XML → ValueError 가 meta.error 로 기록."""
    bad = tmp_path / "weird.hwp"
    bad.write_text(
        '<?xml version="1.0"?><lawml><para>not hwpml</para></lawml>',
        encoding="utf-8",
    )
    r = extract_file(bad)
    assert r.ext == ".hwp"
    assert not r.ok
    assert "expected HWPML" in (r.meta.get("error") or "")


def test_extract_pdf_fake(tmp_path: Path):
    """pypdf 는 빈 PDF 에도 오류 없이 빈 텍스트 반환."""
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    pdf = tmp_path / "blank.pdf"
    with pdf.open("wb") as f:
        writer.write(f)
    r = extract_file(pdf)
    assert r.ext == ".pdf"


def test_extract_docx(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document

    d = Document()
    d.add_heading("제목", level=1)
    d.add_paragraph("본문 단락입니다.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    path = tmp_path / "doc.docx"
    d.save(str(path))
    r = extract_file(path)
    assert r.ok
    assert "본문 단락" in r.text


def test_extract_xlsx(tmp_path: Path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append(["col1", "col2"])
    ws.append(["v1", 42])
    path = tmp_path / "data.xlsx"
    wb.save(str(path))
    r = extract_file(path)
    assert r.ok
    assert "col1" in r.text
    assert "v1" in r.text


def test_extract_directory_skips_hidden(tmp_path: Path):
    (tmp_path / ".hidden.md").write_text("hidden", encoding="utf-8")
    (tmp_path / "visible.md").write_text("visible content", encoding="utf-8")
    results = extract_directory(tmp_path)
    assert all("visible" in r.text or r.path.name == "visible.md" for r in results)
    assert not any(r.path.name == ".hidden.md" for r in results)


def test_extract_directory_filters_unsupported(tmp_path: Path):
    (tmp_path / "a.md").write_text("md", encoding="utf-8")
    (tmp_path / "b.bin").write_bytes(b"\x00")
    results = extract_directory(tmp_path)
    paths = [r.path.name for r in results]
    assert "a.md" in paths
    assert "b.bin" not in paths


def test_max_chars_truncation(tmp_path: Path):
    big = tmp_path / "big.md"
    big.write_text("X" * 10_000, encoding="utf-8")
    r = extract_file(big, max_chars=500)
    assert len(r.text) <= 600
    assert "truncated" in r.text


def test_combined_text_respects_total_limit(tmp_path: Path):
    (tmp_path / "a.md").write_text("A" * 2000, encoding="utf-8")
    (tmp_path / "b.md").write_text("B" * 2000, encoding="utf-8")
    results = extract_directory(tmp_path)
    combined = combined_text(results, max_total_chars=1500)
    assert len(combined) <= 1700
